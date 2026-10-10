import os
from pathlib import Path
import runpy
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import openpyxl
import xlrd

SCRIPT = Path(__file__).resolve().parents[1] / 'prepare-school-data.py'


class PrepareSchoolDataTest(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix='student-web-prepare-', dir=os.environ.get('STUDENT_WEB_TEST_TMP')))
        self.schedule = self.root / 'schedule'
        self.rosters = self.root / 'rosters'
        self.schedule.mkdir()
        self.rosters.mkdir()
        self.output = self.root / 'output.json'
        self.catalog = self.root / 'catalog.json'
        self.output.write_text('previous manifest', encoding='utf-8')
        self.catalog.write_text('previous catalog', encoding='utf-8')

    def run_prepare(self, courses=True):
        sheet = SimpleNamespace(nrows=3, ncols=3, cell_value=lambda row, column: 'Java/周一' if courses and row == 2 and column == 2 else '')
        with patch.object(sys, 'argv', [str(SCRIPT), str(self.output), '--schedule', str(self.schedule), '--rosters', str(self.rosters), '--catalog', str(self.catalog)]), \
                patch.object(xlrd, 'open_workbook', return_value=SimpleNamespace(sheets=lambda: [sheet])):
            runpy.run_path(str(SCRIPT), run_name='__main__')

    def assert_unchanged(self):
        self.assertEqual(self.output.read_text(encoding='utf-8'), 'previous manifest')
        self.assertEqual(self.catalog.read_text(encoding='utf-8'), 'previous catalog')

    def roster(self, populated=True):
        workbook = openpyxl.Workbook()
        if populated:
            workbook.active.append([''] * 8)
            workbook.active.cell(5, 3, '202500000001')
            workbook.active.cell(5, 4, 'Test Student')
            workbook.active.cell(5, 6, '2025级计算机科学与技术')
            workbook.active.cell(5, 8, '2025级计科1班')
        workbook.save(self.rosters / 'roster.xlsx')
        workbook.close()
        (self.schedule / '2025计科1班课表.xls').write_bytes(b'mocked workbook')

    def test_missing_directory_preserves_outputs(self):
        self.schedule = self.root / 'missing'
        with self.assertRaisesRegex(ValueError, '课表目录不存在'):
            self.run_prepare()
        self.assert_unchanged()

    def test_empty_directory_preserves_outputs(self):
        with self.assertRaisesRegex(ValueError, '没有 xlsx 文件'):
            self.run_prepare()
        self.assert_unchanged()

    def test_unparseable_roster_preserves_outputs(self):
        self.roster(populated=False)
        with self.assertRaisesRegex(ValueError, '未识别到有效学生'):
            self.run_prepare()
        self.assert_unchanged()

    def test_insufficient_roster_columns_report_the_file_and_row(self):
        workbook = openpyxl.Workbook()
        workbook.active.cell(5, 1, 'Wrong format')
        workbook.save(self.rosters / 'narrow.xlsx')
        workbook.close()
        (self.schedule / '2025计科1班课表.xls').write_bytes(b'mocked workbook')
        with self.assertRaisesRegex(ValueError, '名单列数不足：narrow.xlsx 第 5 行'):
            self.run_prepare()
        self.assert_unchanged()

    def test_empty_timetable_preserves_outputs(self):
        self.roster()
        with self.assertRaisesRegex(ValueError, '未识别到有效课程'):
            self.run_prepare(courses=False)
        self.assert_unchanged()

    def test_fractional_student_number_is_rejected_without_truncation(self):
        self.roster()
        workbook = openpyxl.load_workbook(self.rosters / 'roster.xlsx')
        workbook.active.cell(5, 3, 202500000001.5)
        workbook.save(self.rosters / 'roster.xlsx')
        workbook.close()
        with self.assertRaisesRegex(ValueError, '学号不能含小数：roster.xlsx 第 5 行'):
            self.run_prepare()
        self.assert_unchanged()

    def test_validated_manifest_contains_its_catalog(self):
        import json
        self.roster()
        self.run_prepare()
        result = json.loads(self.output.read_text(encoding='utf-8'))
        self.assertEqual(len(result['students']), 1)
        self.assertEqual(result['catalog'], json.loads(self.catalog.read_text(encoding='utf-8')))
        self.assertEqual(result['catalog'][0]['classes'][0]['courseCodes'], [result['courses'][0]['code']])


if __name__ == '__main__':
    unittest.main()
